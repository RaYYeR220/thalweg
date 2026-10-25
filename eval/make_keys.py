"""Build the answer keys for the VGC2020 benchmark (reference discharge, its uncertainty, expert answers).

The keys are written to eval/keys/ (git-ignored) and are read only by the scorer, never by the
pipeline or the agent. Nothing is printed except the number of entries written.
"""

import csv
import json
from pathlib import Path

import openpyxl

from thalweg.paths import data_dir

ROOT = data_dir() / "vgc2020" / "VGC2020 (dataset)"
OUT = Path(__file__).parent / "keys" / "vgc.json"


def main():
    wb = openpyxl.load_workbook(ROOT / "Video_Globe_Challenge_2020.xlsx", data_only=True)
    keys = {}
    for n in range(1, 9):
        rows = list(csv.DictReader(open(ROOT / f"Video_{n}" / f"VGC{n}.csv", encoding="latin-1")))
        experts = [float(r["q"]) for r in rows if r.get("q")]
        ws = wb[f"GDH{n}"]
        q_ref = float(ws["C2"].value)
        unc = None
        for col in ("E", "F"):
            hdr = (ws[f"{col}1"].value or "").lower()
            val = ws[f"{col}2"].value
            if "incertitude" in hdr and isinstance(val, (int, float)):
                unc = float(val)
        site = ws["I1"].value if n in (1, 2) else ws["H1"].value
        keys[f"vgc{n}"] = {"q_ref": q_ref, "u_ref_raw": unc, "site": site, "experts": experts,
                           "n_experts": len(experts)}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(keys, indent=1, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {len(keys)} keys to {OUT}")


if __name__ == "__main__":
    main()
